#!/usr/bin/env python3
"""Fill the department's per-variant "Weather Advisory" report template.

One tab per weekly campaign; one row per variant, showing the districts and
mandals that variant covered plus Total Sent / Delivered / Failed (+%). The
"Advisory Variant" column is filled with each variant's forecast text (which
contains the suggested rain dates) from the forecast inputs workbook.
Delivered = Total - Failed (the template's two-bucket view). Geography is
joined from the dissemination roster; phone numbers are never touched.

Reference files (in REFERENCE_DIR, none committed to the repo):
  advisory_report_template.xlsx  the dept's blank template (one tab per week)
  forecast_inputs.xlsx           per-variant forecast text; WAT sheet columns:
                                 week, arm, message, unique_variant_tag
  geography_labels.xlsx          optional canonical Title-Case geography labels
"""
import csv, glob, os, re, sys
from collections import defaultdict
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment

import config
DATA     = config.DATA_DIR
DISSEM   = config.require_dissem()
LABELS   = config.REFERENCE_DIR / "geography_labels.xlsx"
TEMPLATE = config.REFERENCE_DIR / "advisory_report_template.xlsx"
FORECAST = config.REFERENCE_DIR / "forecast_inputs.xlsx"
OUT      = DATA / "Weather_Advisory_Variant_wise_report_FILLED.xlsx"

# >>> EDIT PER RUN <<< one entry per template tab: (exact tab name in the
# template, the week label as used in forecast_inputs.xlsx, that week's pulled CSV).
WEEKS = [
    ("TAB NAME 1", "week1", DATA / "REPLACE_with_week1_pull.csv"),
    ("TAB NAME 2", "week2", DATA / "REPLACE_with_week2_pull.csv"),
    ("TAB NAME 3", "week3", DATA / "REPLACE_with_week3_pull.csv"),
]

def collapse(s): return re.sub(r"\s+", " ", s.strip())
def _k(s): return collapse(s).lower()

# canonical display labels from the original dept file
dd, dm = {}, {}
try:
    ob = openpyxl.load_workbook(LABELS, read_only=True)
    for r in ob["DF By Mandal"].iter_rows(min_row=2, values_only=True):
        if r[0] and r[0] != "TOTAL" and r[1]:
            dd[_k(r[0])] = collapse(r[0]); dm[_k(r[1])] = collapse(r[1])
    ob.close()
except FileNotFoundError:
    pass
def D(s): return dd.get(_k(s), collapse(s).title())
def M(s): return dm.get(_k(s), collapse(s).title())

print("loading geography...")
geo = {}
with open(DISSEM, encoding="utf-8-sig") as f:
    for row in csv.DictReader(f):
        geo[row["ppbno"].strip()] = (D(row["district_name"]), M(row["mandal_name"]))

def vnum(v):
    m = re.search(r"var(\d+)", v, re.I); return int(m.group(1)) if m else 0

# forecast text per (week, variant tag) from the WAT sheet, which carries the
# full advisory text for EVERY variant (all arms). tag == our variant number.
print("loading forecast text (WAT sheet)...")
fc = {}
fwb = openpyxl.load_workbook(FORECAST, data_only=True)
ws_wat = fwb["WAT"]
for r in range(2, ws_wat.max_row + 1):
    wk, msg, tag = ws_wat.cell(r, 1).value, ws_wat.cell(r, 3).value, ws_wat.cell(r, 4).value
    if wk and tag is not None and msg:
        fc[(str(wk).strip().lower(), int(tag))] = collapse(str(msg))
fwb.close()
print(f"  {len(fc)} forecast messages loaded")

def aggregate(path):
    tot = defaultdict(int); failed = defaultdict(int)
    dists = defaultdict(set); mands = defaultdict(set)
    with open(path, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            v = vnum(row["variant"])
            tot[v] += 1
            if row["last_status"] == "failed": failed[v] += 1
            g = geo.get(row["ppbno"].strip())
            if g:
                dists[v].add(g[0]); mands[v].add(g[1])
    return tot, failed, dists, mands

wb = openpyxl.load_workbook(TEMPLATE)
NUM = "#,##0"; PCT = "0.0%"
wrap = Alignment(wrap_text=True, vertical="top")

for tab, week_label, path in WEEKS:
    ws = wb[tab]
    if not path or not Path(path).exists():
        print(f"SKIP {tab!r}: no data yet ({path})")
        continue
    print(f"filling {tab!r} from {Path(path).name} ...")
    tot, failed, dists, mands = aggregate(path)
    # clear any old data rows (row 2 down), keep header row 1
    for r in range(2, ws.max_row + 1):
        for c in range(1, 10): ws.cell(r, c).value = None
    rr = 2; missing = []
    for sl, v in enumerate(sorted(tot), start=1):
        total = tot[v]; fa = failed[v]; deliv = total - fa
        msg = fc.get((week_label, v))
        if not msg: missing.append(v)
        ws.cell(rr, 1).value = sl
        ws.cell(rr, 2).value = msg if msg else f"var{v}"   # advisory text (contains rain dates)
        ws.cell(rr, 2).alignment = wrap
        ws.cell(rr, 3).value = ", ".join(sorted(dists[v]))
        ws.cell(rr, 4).value = ", ".join(sorted(mands[v]))
        ws.cell(rr, 5).value = total
        ws.cell(rr, 6).value = deliv
        ws.cell(rr, 7).value = fa
        ws.cell(rr, 8).value = deliv / total if total else 0
        ws.cell(rr, 9).value = fa / total if total else 0
        for c in (5, 6, 7): ws.cell(rr, c).number_format = NUM
        for c in (8, 9): ws.cell(rr, c).number_format = PCT
        ws.cell(rr, 3).alignment = wrap; ws.cell(rr, 4).alignment = wrap
        rr += 1
    # totals row
    gt = sum(tot.values()); gf = sum(failed.values()); gd = gt - gf
    ws.cell(rr, 2).value = "TOTAL"
    ws.cell(rr, 5).value = gt; ws.cell(rr, 6).value = gd; ws.cell(rr, 7).value = gf
    ws.cell(rr, 8).value = gd / gt if gt else 0; ws.cell(rr, 9).value = gf / gt if gt else 0
    for c in (5, 6, 7): ws.cell(rr, c).number_format = NUM
    for c in (8, 9): ws.cell(rr, c).number_format = PCT
    miss = f"  MISSING forecast for var{missing}" if missing else "  (all variants matched to forecast text)"
    print(f"   {len(tot)} variants, total {gt:,}, failed {gf:,} ({100*gf/gt:.1f}%){miss}")

wb.save(OUT)
print(f"\nSaved -> {OUT}")
