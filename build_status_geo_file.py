#!/usr/bin/env python3
"""Build the final shareable file for a week's pull:
   ppbno, last_status, last_status_timestamp, district, mandal, cluster.
Joins the latest anon pull (ppbno -> status + status time) to the dissemination
geography (ppbno -> district/mandal/cluster, Title-Case display). Phone-free.

Usage: python build_status_geo_file.py [<anon_csv>] [<out_csv>]
Defaults to the most recent farmers_delivery_status_anon_*.csv.
"""
import csv, glob, os, re, sys
from pathlib import Path
import openpyxl

import config
DATA = config.DATA_DIR
DISSEM = config.require_dissem()
LABELS = config.REFERENCE_DIR / "geography_labels.xlsx"

def collapse(s): return re.sub(r"\s+", " ", s.strip())
def _k(s): return collapse(s).lower()

# canonical display labels from the original dept file (Title Case)
dd, dm, dc = {}, {}, {}
try:
    ob = openpyxl.load_workbook(LABELS, read_only=True)
    for r in ob["DF By Mandal"].iter_rows(min_row=2, values_only=True):
        if r[0] and r[0] != "TOTAL" and r[1]:
            dd[_k(r[0])] = collapse(r[0]); dm[_k(r[1])] = collapse(r[1])
    for r in ob["DF By Cluster"].iter_rows(min_row=2, values_only=True):
        if r[0] and r[0] != "TOTAL" and r[2]:
            dc[_k(r[2])] = collapse(r[2])
    ob.close()
except FileNotFoundError:
    pass
def D(s): return dd.get(_k(s), collapse(s).title())
def M(s): return dm.get(_k(s), collapse(s).title())
def C(s): return dc.get(_k(s), collapse(s).title())

print("loading geography (ppbno -> district/mandal/cluster)...")
geo = {}
with open(DISSEM, encoding="utf-8-sig") as f:
    for row in csv.DictReader(f):
        geo[row["ppbno"].strip()] = (D(row["district_name"]), M(row["mandal_name"]), C(row["cluster_name"]))

src = sys.argv[1] if len(sys.argv) > 1 else sorted(
    glob.glob(str(DATA / "farmers_delivery_status_anon_*.csv")), key=os.path.getmtime)[-1]
out = sys.argv[2] if len(sys.argv) > 2 else str(DATA / "week4_jun19_status_timestamp_geo_ppbno.csv")
print(f"source: {Path(src).name}")

cols = ["ppbno", "last_status", "last_status_timestamp", "district", "mandal", "cluster"]
n = 0; unmatched = 0; missing_ts = 0
with open(src, encoding="utf-8-sig", newline="") as fin, \
     open(out, "w", encoding="utf-8-sig", newline="") as fout:
    r = csv.DictReader(fin)
    if "last_status_timestamp" not in r.fieldnames:
        sys.exit("ERROR: source has no last_status_timestamp column — re-pull with the updated script.")
    w = csv.DictWriter(fout, fieldnames=cols)
    w.writeheader()
    for row in r:
        g = geo.get(row["ppbno"].strip())
        if not g:
            unmatched += 1
            dist = mand = clus = ""
        else:
            dist, mand, clus = g
        if not row.get("last_status_timestamp"):
            missing_ts += 1
        w.writerow({
            "ppbno": row["ppbno"],
            "last_status": row["last_status"],
            "last_status_timestamp": row.get("last_status_timestamp", ""),
            "district": dist, "mandal": mand, "cluster": clus,
        })
        n += 1

print(f"\nwrote {n:,} rows -> {out}")
print(f"unmatched ppbno (no geography): {unmatched}")
print(f"rows missing last_status_timestamp: {missing_ts}")
